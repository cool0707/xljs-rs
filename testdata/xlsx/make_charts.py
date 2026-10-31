"""Builds charts.xlsx: Excel charts of each kind for the tests of
ExcelChart (src/api/xlsx_chart.rs and prelude.js). Run with openpyxl:
python3 testdata/xlsx/make_charts.py"""
import datetime, os
from openpyxl import Workbook
from openpyxl.chart import (BarChart, BarChart3D, LineChart, PieChart, DoughnutChart, AreaChart,
                            ScatterChart, BubbleChart, RadarChart, StockChart, Reference, Series)
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.series import DataPoint

wb = Workbook()
ws = wb.active
ws.title = "売上"
rows = [("月", "売上", "原価", "利益率"), ("4月", 1200, 800, 0.33), ("5月", 1500, 900, 0.4), ("6月", 900, 700, 0.22)]
for r in rows:
    ws.append(r)
for r in range(2, 5):
    ws.cell(r, 2).number_format = '#,##0"円"'
    ws.cell(r, 4).number_format = "0%"
cats = Reference(ws, min_col=1, min_row=2, max_row=4)

def at(chart, cell, title=None):
    if title:
        chart.title = title
    ws.add_chart(chart, cell)

# 1. clustered column with axis titles and a fixed color for the first series
c = BarChart(); c.type = "col"; c.grouping = "clustered"
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
c.y_axis.title = "金額"; c.x_axis.title = "月"
c.series[0].graphicalProperties.solidFill = "1F4E79"
at(c, "F2", "月別売上")

# 2. combo: columns and a line on the secondary axis
c = BarChart(); c.type = "col"
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
l = LineChart()
l.add_data(Reference(ws, min_col=4, min_row=1, max_row=4), titles_from_data=True)
l.y_axis.axId = 200; l.y_axis.title = "利益率"; l.y_axis.number_format = "0%"; l.y_axis.crosses = "max"
c += l
at(c, "F20", "売上と利益率")

# 3. horizontal 100% stacked bars
c = BarChart(); c.type = "bar"; c.grouping = "percentStacked"; c.overlap = 100
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
at(c, "F38", "構成比")

# 4. pie with point colors, 5. doughnut
c = PieChart()
c.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
for i, col in enumerate(["C00000", "00B050", "0070C0"]):
    p = DataPoint(idx=i); p.graphicalProperties.solidFill = col; c.series[0].dPt.append(p)
at(c, "P2", "売上の内訳")
c = DoughnutChart()
c.add_data(Reference(ws, min_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
at(c, "P20")

# 6. stacked area, 7. 3-D column, 8. radar (not supported), 9. line without markers
c = AreaChart(); c.grouping = "stacked"
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
at(c, "P38", "推移")
c = BarChart3D()
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
at(c, "Z2", "立体")
c = RadarChart()
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
at(c, "Z20", "レーダー")
c = LineChart()
c.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
c.set_categories(cats)
for s in c.series:
    s.marker.symbol = "none"
c.y_axis.number_format = "#,##0"
at(c, "Z38", "折れ線")

# Points: scatter and bubble
ps = wb.create_sheet("散布")
for r in [("x", "y", "size"), (1, 2, 10), (2, 5, 30), (3, 3, 20)]:
    ps.append(r)
c = ScatterChart(); c.style = 13
s = Series(Reference(ps, min_col=2, min_row=1, max_row=4), Reference(ps, min_col=1, min_row=2, max_row=4), title_from_data=True)
s.marker.symbol = "circle"; s.graphicalProperties.line.noFill = True
c.series.append(s)
ps.add_chart(c, "E2")
c = BubbleChart()
c.series.append(Series(values=Reference(ps, min_col=2, min_row=2, max_row=4), xvalues=Reference(ps, min_col=1, min_row=2, max_row=4),
                       zvalues=Reference(ps, min_col=3, min_row=2, max_row=4), title="泡"))
ps.add_chart(c, "E20")

# Gantt: stacked horizontal bars whose first series (the start dates) is not filled
g = wb.create_sheet("工程")
data = [("作業", "開始日", "日数"), ("設計", datetime.date(2026, 11, 2), 5), ("実装", datetime.date(2026, 11, 9), 12), ("試験", datetime.date(2026, 11, 23), 5)]
for r in data:
    g.append(r)
for r in range(2, 5):
    g.cell(r, 2).number_format = "yyyy/m/d"
def gantt():
    c = BarChart(); c.type = "bar"; c.grouping = "stacked"; c.overlap = 100
    c.add_data(Reference(g, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    c.set_categories(Reference(g, min_col=1, min_row=2, max_row=4))
    c.series[0].graphicalProperties.noFill = True
    c.series[0].graphicalProperties.line.noFill = True
    c.x_axis.scaling.orientation = "maxMin"
    c.y_axis.number_format = "m/d"
    return c
g.add_chart(gantt(), "E2")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "charts.xlsx")
wb.save(out)

# The line chart ("折れ線", chart9) gets a color style as Excel writes for
# "Change Colors" (openpyxl writes none): accent6, accent5, then darker.
import shutil, zipfile
COLORS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cs:colorStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" meth="cycle" id="15"><a:schemeClr val="accent6"/><a:schemeClr val="accent5"/><cs:variation/><cs:variation><a:lumMod val="60000"/></cs:variation></cs:colorStyle>"""
RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2011/relationships/chartColorStyle" Target="colors1.xml"/></Relationships>"""
tmp = out + ".tmp"
with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            data = data.replace(b"</Types>", b'<Override PartName="/xl/charts/colors1.xml" ContentType="application/vnd.ms-office.chartcolorstyle+xml"/></Types>')
        zout.writestr(item, data)
    zout.writestr("xl/charts/colors1.xml", COLORS)
    zout.writestr("xl/charts/_rels/chart9.xml.rels", RELS)
shutil.move(tmp, out)
