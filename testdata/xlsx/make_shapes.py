"""Builds shapes.xlsx: shapes, text boxes, lines, groups and a freeform as
Excel writes them, for the print layout (src/api/xlsx_shape.rs).
openpyxl cannot write shapes, so the drawing part is added to its output.
Run: python3 testdata/xlsx/make_shapes.py"""
import os, re, zipfile, io
from openpyxl import Workbook

here = os.path.dirname(os.path.abspath(__file__))
wb = Workbook()
ws = wb.active
ws.title = "図形"
ws["A1"] = "図形のテスト"
ws["A30"] = "下端"
for c in "ABCDEFGHIJ":
    ws.column_dimensions[c].width = 10
buf = io.BytesIO()
wb.save(buf)

NS = ('xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
STYLE = ('<xdr:style><a:lnRef idx="2"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef>'
         '<a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef>'
         '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></xdr:style>')
CXN_STYLE = ('<xdr:style><a:lnRef idx="1"><a:schemeClr val="accent1"/></a:lnRef><a:fillRef idx="0"><a:schemeClr val="accent1"/></a:fillRef>'
             '<a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef><a:fontRef idx="minor"><a:schemeClr val="tx1"/></a:fontRef></xdr:style>')
ids = iter(range(2, 100))

def anchor(c1, r1, c2, r2, body, client=""):
    m = lambda tag, c, r: f'<xdr:{tag}><xdr:col>{c}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:{tag}>'
    return f'<xdr:twoCellAnchor>{m("from", c1, r1)}{m("to", c2, r2)}{body}<xdr:clientData{client}/></xdr:twoCellAnchor>'

def xfrm(rot=0, flip=""):
    r = f' rot="{rot}"' if rot else ""
    return f'<a:xfrm{r}{flip}><a:off x="0" y="0"/><a:ext cx="1000000" cy="500000"/></a:xfrm>'

def text(paras, anchor="ctr", algn="ctr"):
    ps = "".join(f'<a:p><a:pPr algn="{algn}"/>{p}</a:p>' for p in paras)
    return f'<xdr:txBody><a:bodyPr vertOverflow="clip" rtlCol="0" anchor="{anchor}"/><a:lstStyle/>{ps}</xdr:txBody>'

def run(t, extra="", color=None):
    fill = f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>' if color else ""
    return f'<a:r><a:rPr lang="ja-JP" altLang="en-US" sz="1100"{extra}>{fill}</a:rPr><a:t>{t}</a:t></a:r>'

def sp(prst, name, txt="", sppr_extra="", style=STYLE, av="", rot=0, hidden=False, tx_box=False):
    h = ' hidden="1"' if hidden else ""
    tb = ' txBox="1"' if tx_box else ""
    return (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{next(ids)}" name="{name}"{h}/><xdr:cNvSpPr{tb}/></xdr:nvSpPr>'
            f'<xdr:spPr>{xfrm(rot)}<a:prstGeom prst="{prst}"><a:avLst>{av}</a:avLst></a:prstGeom>{sppr_extra}</xdr:spPr>{style}{txt}</xdr:sp>')

shapes = []
# A text box: white with a black line (as Excel's Insert > Text Box).
shapes.append(anchor(0, 2, 3, 6, sp("rect", "テキスト ボックス 1",
    text([run("テキストボックスの文字"), run("太字で赤", ' b="1"', "FF0000")], anchor="t", algn="l"),
    '<a:solidFill><a:schemeClr val="lt1"/></a:solidFill><a:ln w="9525" cmpd="sng"><a:solidFill><a:schemeClr val="lt1"><a:shade val="50000"/></a:schemeClr></a:solidFill></a:ln>',
    style=STYLE.replace('<a:schemeClr val="lt1"/></a:fontRef>', '<a:schemeClr val="dk1"/></a:fontRef>'), tx_box=True)))
# Default shapes (theme accent1 fill, darker line, white text).
shapes.append(anchor(4, 2, 6, 6, sp("roundRect", "角丸四角形 2", text([run("角丸")]))))
shapes.append(anchor(7, 2, 9, 6, sp("ellipse", "楕円 3", text([run("楕円")]),
    '<a:solidFill><a:srgbClr val="70AD47"/></a:solidFill>')))
shapes.append(anchor(0, 7, 2, 10, sp("rightArrow", "右矢印 4")))
shapes.append(anchor(3, 7, 6, 11, sp("wedgeRoundRectCallout", "吹き出し 5", text([run("ここに注意")]),
    '<a:solidFill><a:srgbClr val="FFF2CC"/></a:solidFill><a:ln w="12700"><a:solidFill><a:srgbClr val="BF9000"/></a:solidFill></a:ln>',
    av='<a:gd name="adj1" fmla="val -40000"/><a:gd name="adj2" fmla="val 90000"/><a:gd name="adj3" fmla="val 16667"/>',
    style=STYLE.replace('<a:schemeClr val="lt1"/></a:fontRef>', '<a:schemeClr val="dk1"/></a:fontRef>'))))
shapes.append(anchor(7, 7, 9, 11, sp("star5", "星 6", "", '<a:solidFill><a:srgbClr val="FFC000"/></a:solidFill>', rot=1200000)))
# A line with an arrowhead, a dashed elbow connector.
shapes.append(anchor(0, 12, 3, 15, f'<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="{next(ids)}" name="直線矢印コネクタ 7"/><xdr:cNvCxnSpPr/></xdr:nvCxnSpPr>'
    f'<xdr:spPr>{xfrm()}<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln w="19050"><a:solidFill><a:srgbClr val="C00000"/></a:solidFill><a:tailEnd type="triangle"/></a:ln></xdr:spPr>{CXN_STYLE}</xdr:cxnSp>'))
shapes.append(anchor(4, 12, 7, 15, f'<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="{next(ids)}" name="カギ線コネクタ 8"/><xdr:cNvCxnSpPr/></xdr:nvCxnSpPr>'
    f'<xdr:spPr>{xfrm(flip=" flipV=\"1\"")}<a:prstGeom prst="bentConnector3"><a:avLst/></a:prstGeom><a:ln w="12700"><a:prstDash val="dash"/><a:headEnd type="oval"/><a:tailEnd type="arrow"/></a:ln></xdr:spPr>{CXN_STYLE}</xdr:cxnSp>'))
# A group of a rectangle and a diamond (child coordinates scaled to the anchor).
g_child = lambda prst, name, x, color: (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{next(ids)}" name="{name}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr><a:xfrm><a:off x="{x}" y="0"/><a:ext cx="400000" cy="400000"/></a:xfrm><a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
    f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill></xdr:spPr>{STYLE}{text([run(name)])}</xdr:sp>')
shapes.append(anchor(0, 16, 4, 20, f'<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="{next(ids)}" name="グループ 9"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr>'
    f'<xdr:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="400000"/><a:chOff x="0" y="0"/><a:chExt cx="1000000" cy="400000"/></a:xfrm></xdr:grpSpPr>'
    + g_child("rect", "四角", 0, "5B9BD5") + g_child("diamond", "ひし形", 600000, "ED7D31") + '</xdr:grpSp>'))
# A freeform (a zigzag line).
shapes.append(anchor(5, 16, 9, 20, f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{next(ids)}" name="フリーフォーム 10"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    f'<xdr:spPr>{xfrm()}<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="0" b="0"/><a:pathLst><a:path w="400" h="100">'
    '<a:moveTo><a:pt x="0" y="100"/></a:moveTo><a:lnTo><a:pt x="100" y="0"/></a:lnTo><a:lnTo><a:pt x="200" y="100"/></a:lnTo>'
    '<a:cubicBezTo><a:pt x="250" y="0"/><a:pt x="350" y="0"/><a:pt x="400" y="100"/></a:cubicBezTo></a:path></a:pathLst></a:custGeom>'
    f'<a:noFill/><a:ln w="28575"><a:solidFill><a:srgbClr val="7030A0"/></a:solidFill></a:ln></xdr:spPr>{STYLE}</xdr:sp>'))
# Not printed: hidden, and "print object" off.
shapes.append(anchor(0, 21, 2, 24, sp("rect", "非表示 11", text([run("見えない")]), hidden=True)))
shapes.append(anchor(3, 21, 5, 24, sp("rect", "印刷しない 12", text([run("印刷しない")])), client=' fPrintsWithSheet="0"'))
# Vertical text.
shapes.append(anchor(6, 21, 7, 27, sp("rect", "縦書き 13", text([run("縦書きの文字")]).replace('anchor="ctr"', 'anchor="ctr" vert="eaVert"'),
    '<a:solidFill><a:srgbClr val="DDEBF7"/></a:solidFill>', style=STYLE.replace('<a:schemeClr val="lt1"/></a:fontRef>', '<a:schemeClr val="dk1"/></a:fontRef>'))))

drawing = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<xdr:wsDr {NS}>' + "".join(shapes) + '</xdr:wsDr>'

src = zipfile.ZipFile(buf)
out = zipfile.ZipFile(os.path.join(here, "shapes.xlsx"), "w", zipfile.ZIP_DEFLATED)
for item in src.infolist():
    data = src.read(item.filename)
    if item.filename == "[Content_Types].xml":
        data = data.replace(b"</Types>", b'<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
    if item.filename == "xl/worksheets/sheet1.xml":
        s = data.decode()
        if "xmlns:r=" not in s:
            s = s.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
        s = s.replace("</worksheet>", '<drawing r:id="rId1"/></worksheet>')
        data = s.encode()
    out.writestr(item, data)
out.writestr("xl/worksheets/_rels/sheet1.xml.rels",
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>')
out.writestr("xl/drawings/drawing1.xml", drawing)
out.close()
